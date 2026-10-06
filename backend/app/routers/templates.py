import json
import re
from pathlib import Path
from typing import List, Optional
import yaml
from docx import Document
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from sqlmodel import Session, select
from app.config import settings
from app.db import get_session
from app.models import TemplateMetaTable, UserTable
from app.auth import get_current_user_optional
from app.schemas import NewTemplateInput, Placeholder, RecommendationRule, TableColumn, Template, TemplateAnalysis
from app.services.template_loader import get_template_by_id, load_templates_from_disk
from app.utils import now_iso
import uuid

router = APIRouter(prefix="/templates", tags=["Templates"])


def sanitize_filename(filename: str) -> str:
    base = Path(filename).name
    clean = re.sub(r"[^\w\d.-]", "_", base).strip("._")
    return clean or "uploaded_template"


def infer_placeholder_meta(key: str) -> Placeholder:
    label = key.replace("_", " ").title()
    is_secret = bool(re.search(r"bank|gst|mobile|account", key, re.IGNORECASE))
    is_official = bool(re.search(r"ref(?:_no)?|approval|date|signatory", key, re.IGNORECASE))
    is_draft = bool(re.search(r"description|objective", key, re.IGNORECASE))
    is_num = bool(re.search(r"amount|qty|cost|total", key, re.IGNORECASE))
    is_date = bool(re.search(r"date", key, re.IGNORECASE))

    f_type = "date" if is_date else "number" if is_num else "longtext" if is_draft else "text"

    return Placeholder(
        key=key,
        label=label,
        type=f_type,
        required=True,
        question=f"What is the {label.lower()}?",
        never_ai=is_secret or is_official,
        masked=is_secret,
        ai_draftable=is_draft and not (is_secret or is_official),
        section="General",
    )


def scan_docx_elements(doc: Document) -> tuple[list[str], dict[str, list[TableColumn]]]:
    raw_texts: list[str] = []
    tables_detected: dict[str, list[TableColumn]] = {}

    def process_table(tbl):
        loop_var = None
        table_key = None
        for row in tbl.rows:
            row_str = " ".join([c.text for c in row.cells])
            raw_texts.append(row_str)
            m = re.search(r"{%\s*(?:tr\s+)?for\s+([a-zA-Z_]\w*)\s+in\s+([a-zA-Z_]\w*)\s*%}", row_str)
            if m:
                loop_var = m.group(1)
                table_key = m.group(2)
                if table_key not in tables_detected:
                    tables_detected[table_key] = []

        if loop_var and table_key:
            seen_cols = set()
            for row in tbl.rows:
                for c in row.cells:
                    col_pat = r"\{\{\s*" + re.escape(loop_var) + r"\.([a-zA-Z_]\w*).*?\}\}"
                    col_matches = re.findall(col_pat, c.text)
                    for col_name in col_matches:
                        if col_name not in seen_cols:
                            seen_cols.add(col_name)
                            col_label = col_name.replace("_", " ").title()
                            is_num = bool(re.search(r"amount|qty|cost|total|price", col_name, re.I))
                            is_secret = bool(re.search(r"bank|gst|mobile|account", col_name, re.I))
                            is_official = bool(re.search(r"ref_no|approval|date|signatory", col_name, re.I))
                            tables_detected[table_key].append(
                                TableColumn(
                                    key=col_name,
                                    label=col_label,
                                    type="number" if is_num else "text",
                                    required=True,
                                    never_ai=is_secret or is_official,
                                    masked=is_secret,
                                )
                            )

    # 1. Body paragraphs
    for p in doc.paragraphs:
        raw_texts.append(p.text)

    # 2. Body tables
    for tbl in doc.tables:
        process_table(tbl)

    # 3. Headers and Footers
    for s in doc.sections:
        for p in s.header.paragraphs:
            raw_texts.append(p.text)
        for tbl in s.header.tables:
            process_table(tbl)
        for p in s.footer.paragraphs:
            raw_texts.append(p.text)
        for tbl in s.footer.tables:
            process_table(tbl)

    return raw_texts, tables_detected


@router.get("", response_model=List[Template])
def list_templates(session: Session = Depends(get_session)):
    return load_templates_from_disk(session)


@router.get("/{template_id}", response_model=Template)
def get_template(template_id: str, session: Session = Depends(get_session)):
    t = get_template_by_id(template_id, session)
    if not t:
        raise HTTPException(status_code=404, detail="Template not found")
    return t


@router.get("/{template_id}/file")
def download_template_file(template_id: str, session: Session = Depends(get_session)):
    t_dir = settings.TEMPLATES_DIR / template_id
    if not t_dir.exists():
        t_dir = settings.TEMPLATES_DIR / template_id.replace("-", "_")
    f_path = t_dir / "template.docx"
    if not f_path.exists():
        raise HTTPException(status_code=404, detail="Template file not found")
    return FileResponse(
        path=str(f_path),
        filename=f"{template_id}.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@router.post("/analyze", response_model=TemplateAnalysis)
async def analyze_template(
    file: UploadFile = File(...),
    schema_file: Optional[UploadFile] = File(None, alias="schema"),
):
    warnings: list[str] = []
    clean_filename = sanitize_filename(file.filename or "template.docx")
    clean_schema_filename = sanitize_filename(schema_file.filename) if schema_file else None

    tmp_path = settings.DATA_DIR / f"temp_{uuid.uuid4().hex}_{clean_filename}"
    detected: list[Placeholder] = []

    try:
        content = await file.read()
        tmp_path.write_bytes(content)
        doc = Document(str(tmp_path))

        raw_texts, tables_detected = scan_docx_elements(doc)
        full_text = "\n".join(raw_texts)

        # 1. Add table placeholders
        for tbl_key, cols in tables_detected.items():
            tbl_label = tbl_key.replace("_", " ").title()
            detected.append(
                Placeholder(
                    key=tbl_key,
                    label=tbl_label,
                    type="table",
                    required=True,
                    question=f"Please provide the {tbl_label.lower()} entries.",
                    never_ai=False,
                    masked=False,
                    ai_draftable=False,
                    section="Details",
                    columns=cols,
                )
            )

        # 2. Extract scalar placeholders {{ key }}
        scalar_matches = re.findall(r"\{\{\s*([a-zA-Z_]\w*).*?\}\}", full_text)
        detected_scalar_keys = [
            k for k in sorted(list(set(scalar_matches)))
            if k not in tables_detected and not k.startswith("loop") and k not in ["for", "endfor", "if", "endif"]
        ]

        for k in detected_scalar_keys:
            detected.append(infer_placeholder_meta(k))

    except Exception as e:
        warnings.append(f"Could not read docx: {str(e)}.")
    finally:
        if tmp_path.exists():
            tmp_path.unlink()

    if not detected:
        warnings.append("No placeholders detected in document.")

    # Optional YAML schema override
    if schema_file:
        try:
            schema_bytes = await schema_file.read()
            schema_dict = yaml.safe_load(schema_bytes)
            if isinstance(schema_dict, dict) and "placeholders" in schema_dict:
                yaml_p_map = {p["key"]: p for p in schema_dict["placeholders"]}
                for p in detected:
                    if p.key in yaml_p_map:
                        yp = yaml_p_map[p.key]
                        p.label = yp.get("label", p.label)
                        p.type = yp.get("type", p.type)
                        p.required = yp.get("required", p.required)
                        p.question = yp.get("question", p.question)
                        p.never_ai = yp.get("never_ai", p.never_ai)
                        p.masked = yp.get("masked", p.masked)
                        p.ai_draftable = yp.get("ai_draftable", p.ai_draftable)
            warnings.append(f"Successfully applied schema from {clean_schema_filename}.")
        except Exception as e:
            warnings.append(f"Could not parse schema YAML: {str(e)}")

    suggested_rules = [
        RecommendationRule(
            id="rule-custom",
            description="Recommended for standard club events.",
            condition="mode != 'Online'",
        )
    ]

    return TemplateAnalysis(
        file_name=clean_filename,
        yaml_file_name=clean_schema_filename,
        detected=detected,
        suggested_rules=suggested_rules,
        warnings=warnings,
    )


@router.post("", response_model=Template)
def create_template(
    payload: NewTemplateInput,
    session: Session = Depends(get_session),
    current_user: Optional[UserTable] = Depends(get_current_user_optional),
):
    if current_user and current_user.role == "member":
        raise HTTPException(
            status_code=403,
            detail="Forbidden: member role cannot create templates. Requires admin or faculty.",
        )
    safe_folder = re.sub(r"[^\w\d-]+", "_", payload.name.lower()).strip("_")
    t_dir = settings.TEMPLATES_DIR / safe_folder
    t_dir.mkdir(parents=True, exist_ok=True)

    # Save meta.yaml
    meta = {
        "id": safe_folder.replace("_", "-"),
        "name": payload.name,
        "short_name": payload.name,
        "description": payload.description,
        "authority": payload.authority,
        "ref_category": payload.ref_category,
        "file_name": payload.file_name,
        "version": 1,
    }
    (t_dir / "meta.yaml").write_text(yaml.dump(meta, sort_keys=False), encoding="utf-8")

    # Save schema.yaml
    schema = {
        "placeholders": [p.model_dump(by_alias=False) for p in payload.placeholders]
    }
    (t_dir / "schema.yaml").write_text(yaml.dump(schema, sort_keys=False), encoding="utf-8")

    # Save rules.yaml
    rules = {
        "rules": [r.model_dump(by_alias=False) for r in payload.rules]
    }
    (t_dir / "rules.yaml").write_text(yaml.dump(rules, sort_keys=False), encoding="utf-8")

    # If template.docx doesn't exist, generate a basic one
    docx_file = t_dir / "template.docx"
    if not docx_file.exists():
        doc = Document()
        doc.add_heading(payload.name, level=1)
        for p in payload.placeholders:
            doc.add_paragraph(f"{p.label}: {{{{ {p.key} }}}}")
        doc.save(str(docx_file))

    # Sync to DB
    all_tpls = load_templates_from_disk(session)
    created = next((t for t in all_tpls if t.id == meta["id"]), None)
    if not created:
        raise HTTPException(status_code=500, detail="Failed to initialize template")
    return created


@router.patch("/{template_id}", response_model=Template)
@router.put("/{template_id}", response_model=Template)
def update_template(
    template_id: str,
    payload: dict,
    session: Session = Depends(get_session),
    current_user: Optional[UserTable] = Depends(get_current_user_optional),
):
    if current_user and current_user.role == "member":
        raise HTTPException(
            status_code=403,
            detail="Forbidden: member role cannot edit templates. Requires admin or faculty.",
        )
    t_dir = settings.TEMPLATES_DIR / template_id
    if not t_dir.exists():
        t_dir = settings.TEMPLATES_DIR / template_id.replace("-", "_")

    if not t_dir.exists():
        raise HTTPException(status_code=404, detail="Template folder not found")

    if "name" in payload or "description" in payload or "authority" in payload or "refCategory" in payload:
        meta_file = t_dir / "meta.yaml"
        meta = yaml.safe_load(meta_file.read_text(encoding="utf-8")) if meta_file.exists() else {}
        if "name" in payload:
            meta["name"] = payload["name"]
        if "description" in payload:
            meta["description"] = payload["description"]
        if "authority" in payload:
            meta["authority"] = payload["authority"]
        if "refCategory" in payload:
            meta["ref_category"] = payload["refCategory"]
        meta_file.write_text(yaml.dump(meta, sort_keys=False), encoding="utf-8")

    if "placeholders" in payload:
        schema_file = t_dir / "schema.yaml"
        schema_file.write_text(
            yaml.dump({"placeholders": payload["placeholders"]}, sort_keys=False),
            encoding="utf-8",
        )

    if "rules" in payload:
        rules_file = t_dir / "rules.yaml"
        rules_file.write_text(
            yaml.dump({"rules": payload["rules"]}, sort_keys=False),
            encoding="utf-8",
        )

    # Update DB record
    db_meta = session.exec(select(TemplateMetaTable).where(TemplateMetaTable.id == template_id)).first()
    if db_meta:
        db_meta.updated_at = now_iso()
        session.add(db_meta)
        session.commit()

    all_tpls = load_templates_from_disk(session)
    updated = next((t for t in all_tpls if t.id == template_id), None)
    if not updated:
        raise HTTPException(status_code=404, detail="Template not found after update")
    return updated
