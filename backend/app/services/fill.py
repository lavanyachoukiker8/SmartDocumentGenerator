import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List
import yaml
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docxtpl import DocxTemplate
from app.config import settings


def num_to_words_indian(num: Any) -> str:
    """Converts a number to Indian numbering currency words (Lakhs, Crores)."""
    try:
        n = int(round(float(num)))
    except (ValueError, TypeError):
        return str(num) if num else ""

    if n == 0:
        return "Zero Rupees Only"

    units = [
        "", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine", "Ten",
        "Eleven", "Twelve", "Thirteen", "Fourteen", "Fifteen", "Sixteen", "Seventeen", "Eighteen", "Nineteen"
    ]
    tens = ["", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy", "Eighty", "Ninety"]

    def two_digits(val: int) -> str:
        if val < 20:
            return units[val]
        return (tens[val // 10] + (" " + units[val % 10] if val % 10 != 0 else "")).strip()

    def three_digits(val: int) -> str:
        h = val // 100
        rem = val % 100
        res = ""
        if h > 0:
            res += units[h] + " Hundred"
        if rem > 0:
            res += (" " if res else "") + two_digits(rem)
        return res

    parts = []
    crore = n // 10000000
    n %= 10000000
    if crore > 0:
        parts.append(two_digits(crore) + " Crore")

    lakh = n // 100000
    n %= 100000
    if lakh > 0:
        parts.append(two_digits(lakh) + " Lakh")

    thousand = n // 1000
    n %= 1000
    if thousand > 0:
        parts.append(two_digits(thousand) + " Thousand")

    if n > 0:
        parts.append(three_digits(n))

    return (" ".join(parts) + " Rupees Only").strip()


amount_in_words = num_to_words_indian



def date_dmy(val: Any) -> str:
    """Formats date to DD/MM/YYYY."""
    if not val or val == "[TO BE FILLED]":
        return "[TO BE FILLED]"
    s = str(val).strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        y, mo, d = m.groups()
        return f"{int(d):02d}/{int(mo):02d}/{y}"
    return s


def date_range(start_val: Any, end_val: Any = None) -> str:
    """Formats a date or date range."""
    s = date_dmy(start_val)
    if not end_val or str(end_val).strip() == str(start_val).strip():
        return s
    e = date_dmy(end_val)
    return f"{s} to {e}"


def time_12h(val: Any) -> str:
    """Converts 24h time to 12h AM/PM time."""
    if not val or val == "[TO BE FILLED]":
        return "[TO BE FILLED]"
    s = str(val).strip()
    m = re.match(r"^(\d{1,2}):(\d{2})", s)
    if m:
        h, mn = int(m.group(1)), m.group(2)
        ampm = "AM" if h < 12 else "PM"
        h12 = h % 12
        if h12 == 0:
            h12 = 12
        return f"{h12:02d}:{mn} {ampm}"
    return s


def compute_schema_totals(template_id: str, context: Dict[str, Any]) -> None:
    """Computes totals for tables and summaries based on template."""
    # 1. Bill Certificate Items Table
    if "items" in context and isinstance(context["items"], list):
        total_items_cost = 0.0
        for row in context["items"]:
            try:
                qty = float(row.get("qty", 0) or 0)
                unit_cost = float(row.get("unit_cost", 0) or 0)
                if not row.get("total"):
                    row["total"] = round(qty * unit_cost, 2)
                total_items_cost += float(row.get("total", 0) or 0)
            except (ValueError, TypeError):
                pass
        context["total_amount"] = round(total_items_cost, 2)
        taxes = float(context.get("taxes", 0) or 0)
        context["grand_total"] = round(total_items_cost + taxes, 2)

    # 2. Bill Summary Reimbursements Table
    if "reimbursements" in context and isinstance(context["reimbursements"], list):
        tot_prizes = 0.0
        tot_other = 0.0
        tot_grand = 0.0
        for r in context["reimbursements"]:
            try:
                pm = float(r.get("prize_money", 0) or 0)
                ot = float(r.get("other", 0) or 0)
                if not r.get("total"):
                    r["total"] = round(pm + ot, 2)
                tot_prizes += pm
                tot_other += ot
                tot_grand += float(r.get("total", 0) or 0)
            except (ValueError, TypeError):
                pass
        context["total_prizes"] = round(tot_prizes, 2)
        context["total_other"] = round(tot_other, 2)
        context["grand_total"] = round(tot_grand, 2)

    # 3. Bill Summary Other Bank Accounts
    if "other_accounts" in context and isinstance(context["other_accounts"], list):
        tot_oth = 0.0
        for acc in context["other_accounts"]:
            try:
                tot_oth += float(acc.get("amount", 0) or 0)
            except (ValueError, TypeError):
                pass
        context["other_accounts_total"] = round(tot_oth, 2)


def highlight_unfilled_runs(doc: Document) -> None:
    """Highlights runs containing '[TO BE FILLED]' with yellow color."""
    def process_p(p):
        for run in p.runs:
            if "[TO BE FILLED" in run.text:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    for p in doc.paragraphs:
        process_p(p)

    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    process_p(p)


import jinja2


def get_jinja_env() -> jinja2.Environment:
    env = jinja2.Environment()
    env.filters["date_dmy"] = date_dmy
    env.filters["date_range"] = date_range
    env.filters["time_12h"] = time_12h
    env.filters["amount_in_words"] = num_to_words_indian
    return env


def render_docx_template(template_id: str, values: Dict[str, Any], output_path: Path) -> Path:
    """
    Renders values into docxtpl DocxTemplate with custom Jinja filters,
    computed totals, and post-processed yellow highlighting for unfilled fields.
    """
    t_dir = settings.TEMPLATES_DIR / template_id
    if not t_dir.exists():
        t_dir = settings.TEMPLATES_DIR / template_id.replace("-", "_")

    template_file = t_dir / "template.docx"
    if not template_file.exists():
        raise FileNotFoundError(f"Template docx file not found for {template_id}")

    schema_file = t_dir / "schema.yaml"
    required_keys = set()
    if schema_file.exists():
        try:
            s_data = yaml.safe_load(schema_file.read_text(encoding="utf-8"))
            for p in s_data.get("placeholders", []):
                if p.get("required"):
                    required_keys.add(p["key"])
        except Exception:
            pass

    doc = DocxTemplate(str(template_file))

    # Preprocess context
    context: Dict[str, Any] = {}
    for k, v in values.items():
        if v is None or (isinstance(v, str) and not v.strip()):
            context[k] = "[TO BE FILLED]"
        else:
            context[k] = v

    # Inject [TO BE FILLED] for missing schema fields
    for rk in required_keys:
        if rk not in context or context[rk] is None or (isinstance(context[rk], str) and not context[rk].strip()):
            context[rk] = "[TO BE FILLED]"

    # Compute schema totals
    compute_schema_totals(template_id, context)

    # Render docxtpl template
    doc.render(context, jinja_env=get_jinja_env())

    # Post-process: Highlight any [TO BE FILLED] runs yellow
    highlight_unfilled_runs(doc.docx)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path
