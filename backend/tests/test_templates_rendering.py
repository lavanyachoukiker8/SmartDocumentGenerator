import pytest
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from app.services.fill import (
    amount_in_words,
    date_dmy,
    date_range,
    num_to_words_indian,
    render_docx_template,
    time_12h,
)
from app.services.export import convert_docx_to_pdf

OUTPUT_DIR = Path("backend/outputs/test_rendered")


@pytest.fixture(autouse=True)
def ensure_output_dir():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def test_jinja_filters():
    # date_dmy
    assert date_dmy("2026-10-14") == "14/10/2026"
    assert date_dmy("2026-09-08T00:00:00") == "08/09/2026"
    assert date_dmy("") == "[TO BE FILLED]"

    # date_range
    assert date_range("2026-10-14", "2026-10-15") == "14/10/2026 to 15/10/2026"
    assert date_range("2026-10-14", "2026-10-14") == "14/10/2026"

    # time_12h
    assert time_12h("18:00") == "06:00 PM"
    assert time_12h("09:30") == "09:30 AM"

    # amount_in_words (Indian numbering)
    assert "Twelve Thousand" in num_to_words_indian(12500)
    assert "Rupees Only" in num_to_words_indian(12500)
    assert num_to_words_indian(100000) == "One Lakh Rupees Only"
    assert num_to_words_indian(5000000) == "Fifty Lakh Rupees Only"


def test_render_and_export_room_permission():
    room_data = {
        "ref_no": "ACM/26-27/ROOM/009",
        "letter_date": "2026-10-14",
        "submitted_to": "HoD(CSE Dept.)",
        "subject": "Requesting Permission to Conduct ACM Executives Meet in Seminar Hall 402, CSE Department.",
        "event_title": "ACM Executives Meet",
        "venue": "Seminar Hall 402 (4th Floor)",
        "start_date": "2026-10-14",
        "end_date": "2026-10-14",
        "time_from": "18:00",
        "time_to": "19:30",
        "description": "ACM NIT Surat is organizing an Executive Meet.",
        "objective": "To provide guidance to the ACM executive team on organizing events.",
        "participants": "150+",
        "mode": "Offline",
        "organizers": "ACM Core 2026-27",
        "faculty_coordinator": "Dr. Sankita J. Patel",
        "contact_1_name": "Ansh Gupta",
        "contact_1_designation": "Chairperson, ACM",
        "contact_1_admn": "I24AI005",
        "contact_1_branch": "Artificial Intelligence",
        "contact_1_mobile": "",
        "contact_2_name": "Sunny Bodiwala",
        "contact_2_designation": "Tech Assistant, Seminar Hall",
        "contact_2_branch": "Computer Science & Engineering",
        "contact_2_mobile": "",
        "contact_3_name": "Shri. Rakesh P. Gohil",
        "contact_3_designation": "Faculty in-charge, Seminar Hall",
        "contact_3_branch": "Computer Science & Engineering",
        "contact_4_name": "Dr. Sankita J. Patel",
        "contact_4_designation": "Chairman, ACM",
        "contact_4_branch": "Computer Science & Engineering",
    }
    docx_path = render_docx_template("room_permission", room_data, OUTPUT_DIR / "test_room.docx")
    assert docx_path.exists()

    # Convert to PDF
    pdf_path = convert_docx_to_pdf(docx_path, OUTPUT_DIR)
    assert pdf_path.exists()
    assert pdf_path.stat().st_size > 10000

    # Assert logos exist in header
    reader = PdfReader(pdf_path)
    assert len(reader.pages) >= 1
    p0 = reader.pages[0]
    assert len(p0.images) >= 2, "SVNIT and ACM logos must exist in the header"

    text = p0.extract_text()
    assert "ACM/26-27/ROOM/009" in text
    assert "SCHEDULE OF EVENT" in text
    assert "We request permission to use" in text


def test_render_and_export_bill_certificate():
    bill_cert_data = {
        "chapter_name": "ACM",
        "chapter_code": "6/52",
        "approval_note_no": "ACM/2026-27/SIH/001",
        "approval_date": "2026-09-08",
        "bill_no": "1059",
        "bill_date": "2026-09-08",
        "bill_amount": 15420.0,
        "supplier_name": "Reliable Stationery & Tech Hub",
        "purchase_order": "NiL",
        "head_of_account": "6/52",
        "advance_drawn": "NiL",
        "certifying_person": "Dr. Sankita J. Patel",
        "items": [
            {"sr_no": 1, "item_name": "Networking Cables & Crimping Tools", "qty": 5, "unit_cost": 450.0},
            {"sr_no": 2, "item_name": "Event Banners & Lanyards", "qty": 120, "unit_cost": 85.0},
            {"sr_no": 3, "item_name": "Technical Certificates & Stationery", "qty": 150, "unit_cost": 19.8},
        ],
        "taxes": 0.0,
        "supplier_gst": "24AAAAA0000A1Z5",
        "other_taxes": "-",
        "invoice_no": "1059",
        "invoice_date": "2026-09-08",
        "payee_name": "Reliable Stationery & Tech Hub",
        "account_number": "50100234561234",
        "account_holder": "Reliable Stationery",
        "bank_name": "State Bank of India",
        "bank_branch": "SVNIT Branch",
        "ifsc": "SBIN0003320",
        "stock_office": "Student Council",
        "stock_register_type": "Recurring Register",
        "stock_serial_no": "14",
        "stock_page_no": "32",
        "stock_date": "2026-09-08",
        "faculty_chairman_name": "Dr. Sankita J. Patel",
        "faculty_chairman_designation": "Associate Professor & Faculty Chairman",
        "department": "CSE",
    }
    docx_path = render_docx_template("bill_certificate", bill_cert_data, OUTPUT_DIR / "test_bill_cert.docx")
    assert docx_path.exists()

    pdf_path = convert_docx_to_pdf(docx_path, OUTPUT_DIR)
    assert pdf_path.exists()

    reader = PdfReader(pdf_path)
    assert len(reader.pages) >= 2
    assert len(reader.pages[0].images) >= 2, "Logos must exist in page 1 header"

    full_text = "".join(p.extract_text() for p in reader.pages)
    assert "DEAN STUDENT WELFARE" in full_text
    assert "FORM BC-R" in full_text
    assert "Dr. Sankita J. Patel" in full_text
    assert "personally satisfied that the goods" in full_text
    assert "(FOR USE IN ACCOUNT SECTION)" in full_text


def test_render_and_export_bill_summary():
    bill_sum_data = {
        "letter_date": "2026-09-10",
        "event_title": "SIH Internal Hackathon",
        "reimbursements": [
            {"sr_no": 1, "name": "Team ByteForce", "account_details": "HDFC Bank", "paid_to": "Ansh Gupta", "prize_money": 10000.0, "other": 0.0},
            {"sr_no": 2, "name": "Refreshments", "account_details": "SBI Bank", "paid_to": "Arshad Khatib", "prize_money": 0.0, "other": 3500.0},
        ],
        "other_accounts": [
            {"sr_no": 1, "name": "ACM Chapter", "account_no": "50100234561234", "ifsc": "SBIN0003320", "amount": 13500.0}
        ],
        "student_chairperson_name": "Ansh Gupta",
        "student_secretary_name": "Arshad Khatib",
        "faculty_chairman_name": "Dr. Sankita Patel",
    }
    docx_path = render_docx_template("bill_summary", bill_sum_data, OUTPUT_DIR / "test_bill_sum.docx")
    assert docx_path.exists()

    pdf_path = convert_docx_to_pdf(docx_path, OUTPUT_DIR)
    assert pdf_path.exists()

    reader = PdfReader(pdf_path)
    assert len(reader.pages) >= 2
    assert len(reader.pages[0].images) >= 2

    full_text = "".join(p.extract_text() for p in reader.pages)
    assert "BILL SUMMARY" in full_text
    assert "SIH Internal Hackathon" in full_text
    assert "Other Bank Account Details" in full_text
    assert "Ansh Gupta" in full_text
    assert "Dr. Sankita Patel" in full_text


def test_unfilled_fields_are_highlighted_yellow():
    # Missing required fields like venue, start_date, subject
    sparse_data = {"event_title": "Partial Meeting"}
    docx_path = render_docx_template("room_permission", sparse_data, OUTPUT_DIR / "test_sparse.docx")
    doc = Document(str(docx_path))

    highlighted = []
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.highlight_color == WD_COLOR_INDEX.YELLOW:
                highlighted.append(r.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.font.highlight_color == WD_COLOR_INDEX.YELLOW:
                            highlighted.append(r.text)

    assert len(highlighted) > 0, "Unfilled fields must have yellow highlight"
    assert any("[TO BE FILLED]" in t for t in highlighted)
