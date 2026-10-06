import concurrent.futures
import io
import shutil
import tempfile
from pathlib import Path
from docx import Document
from fastapi.testclient import TestClient
from sqlmodel import Session, select
from app.main import app
from app.config import settings
from app.db import engine
from app.models import ClubTable, CounterTable, DocumentTable, EventTable
from app.services.reference import reserve_reference_sequence, reserve_document_reference
from app.services.rules import eval_rule_safe, normalize_condition, validate_rule_condition
from app.services.template_loader import load_templates_from_disk

client = TestClient(app)


def test_reference_concurrency_and_monotonicity():
    # 10 parallel requests to reserve sequence for same category get distinct monotonic numbers with no skips or duplicates
    category = "CONCURRENT_TEST"
    year = "26-27"

    # Reset any counter for this test category
    with Session(engine) as session:
        session.exec(select(CounterTable).where(CounterTable.category == category))
        for row in session.exec(select(CounterTable).where(CounterTable.category == category)).all():
            session.delete(row)
        session.commit()

    def worker(_):
        with Session(engine) as s:
            return reserve_reference_sequence(category, year, s)

    with concurrent.futures.ThreadPoolExecutor(max_workers=10) as executor:
        results = list(executor.map(worker, range(10)))

    # Must be 10 unique sequences
    assert len(results) == 10
    assert len(set(results)) == 10
    # Must be strictly monotonic from 1 to 10
    sorted_res = sorted(results)
    assert sorted_res == list(range(1, 11))


def test_exactly_one_reference_per_document_on_approval():
    # Create event
    e_res = client.post(
        "/api/events/from-form",
        json={"title": "Ref Test Event", "category": "meeting", "startDate": "2026-11-20"},
    )
    assert e_res.status_code == 200
    e_id = e_res.json()["id"]

    # Generate document
    gen_res = client.post(
        f"/api/events/{e_id}/generate",
        json={"templateIds": ["room-permission"], "options": {"allowPlaceholders": True}},
    )
    assert gen_res.status_code == 200
    doc_id = gen_res.json()[0]["id"]

    # Initial draft should NOT have reserved a ref_no
    doc_data = client.get(f"/api/documents/{doc_id}").json()
    init_ref = doc_data["values"].get("ref_no")
    assert init_ref is None or init_ref == "" or init_ref == "[TO BE FILLED]"

    # Approve document: reserves exactly one ref_no
    appr1 = client.patch(f"/api/documents/{doc_id}/status", json={"status": "approved"}).json()
    ref1 = appr1["values"].get("ref_no")
    assert ref1 is not None and "ROOM" in ref1

    # Approving again MUST NOT allocate a new reference number
    appr2 = client.patch(f"/api/documents/{doc_id}/status", json={"status": "approved"}).json()
    ref2 = appr2["values"].get("ref_no")
    assert ref2 == ref1


def test_rules_evaluator_and_all_rules_evaluated():
    # 1. Condition evaluator: support 'is present' and 'is empty'
    ctx = {"venue": "Seminar Hall", "mode": "Offline", "notes": None, "empty_str": ""}
    assert eval_rule_safe("venue is present", ctx) is True
    assert eval_rule_safe("notes is empty", ctx) is True
    assert eval_rule_safe("empty_str is empty", ctx) is True
    assert eval_rule_safe("venue is empty", ctx) is False
    assert eval_rule_safe("venue != None and mode == 'Offline'", ctx) is True

    # 2. Bad expressions return False without crashing, and validate_rule_condition returns error
    matched = eval_rule_safe("syntax error ??? === 123", ctx)
    assert matched is False
    err = validate_rule_condition("syntax error ??? === 123")
    assert err is not None

    # 3. Recommendations evaluate ALL rules and list all reasons
    e_res = client.post(
        "/api/events/from-form",
        json={
            "title": "Hackathon 2026",
            "category": "hackathon",
            "startDate": "2026-12-01",
            "venue": "Auditorium",
            "mode": "Offline",
        },
    )
    e_id = e_res.json()["id"]
    rec_res = client.get(f"/api/events/{e_id}/recommendations")
    assert rec_res.status_code == 200
    recs = rec_res.json()
    # Find room permission recommendation
    room_rec = next((r for r in recs if r["templateId"] == "room-permission"), None)
    assert room_rec is not None
    assert room_rec["recommended"] is True
    assert len(room_rec["reason"]) > 0


def test_analyze_template_and_type_inference():
    # Create docx with header, footer, body paragraphs, and table with {%tr for row in items %}
    doc = Document()
    s = doc.sections[0]
    s.header.paragraphs[0].text = "Header note {{ header_ref }}"
    s.footer.paragraphs[0].text = "Footer info {{ footer_date }}"

    doc.add_paragraph("Event: {{ event_title }} with {{ total_amount }}")
    doc.add_paragraph("Approval: {{ approval_date }} and {{ bank_account }}")

    tbl = doc.add_table(rows=2, cols=3)
    # Loop definition row
    tbl.rows[0].cells[0].text = "{%tr for item in items %}"
    tbl.rows[0].cells[1].text = ""
    tbl.rows[0].cells[2].text = ""
    # Data row
    tbl.rows[1].cells[0].text = "{{ item.item_name }}"
    tbl.rows[1].cells[1].text = "{{ item.item_qty }}"
    tbl.rows[1].cells[2].text = "{{ item.item_cost }}"

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)

    res = client.post(
        "/api/templates/analyze",
        files={"file": ("custom_template.docx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res.status_code == 200
    data = res.json()
    keys = {p["key"]: p for p in data["detected"]}

    # 1. Table placeholder 'items' detected with columns
    assert "items" in keys
    assert keys["items"]["type"] == "table"
    col_map = {c["key"]: c for c in keys["items"]["columns"]}
    assert "item_name" in col_map
    assert col_map["item_qty"]["type"] == "number"
    assert col_map["item_cost"]["type"] == "number"

    # 2. Scalar inferred types and flags
    assert keys["header_ref"]["neverAi"] is True
    assert keys["footer_date"]["type"] == "date"
    assert keys["total_amount"]["type"] == "number"
    assert keys["approval_date"]["neverAi"] is True
    assert keys["bank_account"]["masked"] is True
    assert keys["bank_account"]["neverAi"] is True

    # 3. Never invent fields not in doc
    assert "venue" not in keys
    assert "faculty_coordinator" not in keys


def test_template_loader_schema_errors_reported():
    # Create a temporary template folder without schema.yaml
    bad_dir = settings.TEMPLATES_DIR / "temp_broken_template"
    bad_dir.mkdir(parents=True, exist_ok=True)
    try:
        # Only create template.docx
        doc = Document()
        doc.save(str(bad_dir / "template.docx"))

        with Session(engine) as session:
            templates = load_templates_from_disk(session)
            broken = next((t for t in templates if t.id == "temp-broken-template"), None)
            assert broken is not None
            assert broken.error is not None
            assert "Missing schema.yaml" in broken.error

        # Now write corrupt schema.yaml
        (bad_dir / "schema.yaml").write_text("invalid: [unclosed yaml", encoding="utf-8")
        with Session(engine) as session:
            templates = load_templates_from_disk(session)
            broken = next((t for t in templates if t.id == "temp-broken-template"), None)
            assert broken is not None
            assert broken.error is not None
            assert "Invalid YAML" in broken.error
    finally:
        if bad_dir.exists():
            shutil.rmtree(bad_dir)
