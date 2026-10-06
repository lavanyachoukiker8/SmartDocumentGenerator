import io
import json
import os
import shutil
import tempfile
from pathlib import Path
from docx import Document
from fastapi.testclient import TestClient
from app.main import app
from app.config import settings, get_or_create_fernet_key
from app.crypto import encrypt_field, decrypt_field, mask_value
from app.models import DocumentTable, DocumentVersionTable, EventFieldTable, EventTable
from app.db import engine
from sqlmodel import Session, select

client = TestClient(app)


def test_fernet_key_auto_generation():
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)
        key1 = get_or_create_fernet_key(tmp_path, None)
        assert key1 is not None and len(key1) > 20
        key_file = tmp_path / ".fernet_key"
        assert key_file.exists()
        assert key_file.read_text(encoding="utf-8").strip() == key1

        # Second read uses existing key
        key2 = get_or_create_fernet_key(tmp_path, None)
        assert key2 == key1


def test_masked_fields_encrypted_in_db_and_reveal():
    # Create an event with a masked field (e.g. bank_account)
    event_payload = {
        "title": "Security Test Event",
        "category": "workshop",
        "startDate": "2026-11-01",
        "venue": "Lab 1",
    }
    create_res = client.post("/api/events/from-form", json=event_payload)
    assert create_res.status_code == 200
    event_id = create_res.json()["id"]

    # Generate bill_certificate document
    gen_res = client.post(
        f"/api/events/{event_id}/generate",
        json={"templateIds": ["bill-certificate"], "options": {"allowPlaceholders": True}},
    )
    assert gen_res.status_code == 200
    docs = gen_res.json()
    doc_id = docs[0]["id"]

    # Save document with a sensitive/masked field
    secret_val = "987654321098"
    save_payload = {
        "values": {
            "event_title": "Security Test Event",
            "account_number": secret_val,
        },
        "note": "Adding sensitive details",
    }
    save_res = client.put(f"/api/documents/{doc_id}", json=save_payload)
    assert save_res.status_code == 200
    doc_data = save_res.json()["document"]

    # API response for document must show masked preview, not plaintext secret
    assert "987654321098" not in json.dumps(doc_data["values"])

    # Directly inspect database: must be encrypted ciphertext
    with Session(engine) as session:
        doc_in_db = session.exec(select(DocumentTable).where(DocumentTable.id == doc_id)).first()
        vals = json.loads(doc_in_db.values_json)
        assert "account_number" in vals
        assert vals["account_number"] != secret_val
        # Must decrypt back to original
        assert decrypt_field(vals["account_number"]) == secret_val

        # Version table must also store encrypted
        ver_in_db = session.exec(
            select(DocumentVersionTable)
            .where(DocumentVersionTable.document_id == doc_id)
            .where(DocumentVersionTable.version == doc_in_db.current_version)
        ).first()
        ver_vals = json.loads(ver_in_db.values_json)
        assert ver_vals["account_number"] != secret_val
        assert decrypt_field(ver_vals["account_number"]) == secret_val

    # Reveal endpoint must return decrypted secret
    reveal_res = client.post(
        f"/api/documents/{doc_id}/reveal",
        json={"field_key": "account_number"},
    )
    assert reveal_res.status_code == 200
    assert reveal_res.json()["value"] == secret_val
    assert reveal_res.json()["hasValue"] is True

    # GET reveal endpoint variant
    get_reveal_res = client.get(f"/api/documents/{doc_id}/reveal/account_number")
    assert get_reveal_res.status_code == 200
    assert get_reveal_res.json()["value"] == secret_val


def test_auth_login_and_roles():
    # 1. Invalid login
    res_bad = client.post("/api/auth/login", json={"username": "admin", "password": "wrongpassword"})
    assert res_bad.status_code == 401

    # 2. Member login
    res_member = client.post("/api/auth/login", json={"username": "member", "password": "member123"})
    assert res_member.status_code == 200
    member_token = res_member.json()["access_token"]
    assert res_member.json()["user"]["role"] == "member"

    # 3. Faculty login
    res_faculty = client.post("/api/auth/login", json={"username": "faculty", "password": "faculty123"})
    assert res_faculty.status_code == 200
    faculty_token = res_faculty.json()["access_token"]
    assert res_faculty.json()["user"]["role"] == "faculty"

    # 4. Admin login
    res_admin = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    assert res_admin.status_code == 200
    admin_token = res_admin.json()["access_token"]
    assert res_admin.json()["user"]["role"] == "admin"

    # 5. /api/auth/me
    me_res = client.get("/api/auth/me", headers={"Authorization": f"Bearer {admin_token}"})
    assert me_res.status_code == 200
    assert me_res.json()["username"] == "admin"

    # 6. Role enforcement: member cannot approve document
    # Create test document
    e_res = client.post(
        "/api/events/from-form",
        json={"title": "Role Test Event", "category": "meeting", "startDate": "2026-11-05"},
    )
    e_id = e_res.json()["id"]
    g_res = client.post(
        f"/api/events/{e_id}/generate",
        json={"templateIds": ["room-permission"], "options": {"allowPlaceholders": True}},
    )
    d_id = g_res.json()[0]["id"]

    # Member attempt to approve -> 403 Forbidden
    member_patch = client.patch(
        f"/api/documents/{d_id}/status",
        json={"status": "approved"},
        headers={"Authorization": f"Bearer {member_token}"},
    )
    assert member_patch.status_code == 403

    # Faculty approval -> 200 OK
    faculty_patch = client.patch(
        f"/api/documents/{d_id}/status",
        json={"status": "approved"},
        headers={"Authorization": f"Bearer {faculty_token}"},
    )
    assert faculty_patch.status_code == 200
    assert faculty_patch.json()["status"] == "approved"

    # Member attempt to create template -> 403 Forbidden
    tpl_payload = {
        "name": "Member Template",
        "description": "desc",
        "authority": "auth",
        "refCategory": "GEN",
        "fileName": "test.docx",
        "placeholders": [],
        "rules": [],
    }
    member_tpl = client.post(
        "/api/templates",
        json=tpl_payload,
        headers={"Authorization": f"Bearer {member_token}"},
    )
    assert member_tpl.status_code == 403


def test_sanitize_filename_in_upload():
    # Build minimal docx in memory
    doc = Document()
    doc.add_paragraph("Hello {{ test_placeholder }}")
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)

    # Malicious filename attempting directory traversal
    traversal_name = "../../evil_shell.docx"
    response = client.post(
        "/api/templates/analyze",
        files={"file": (traversal_name, bio.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    res_data = response.json()
    assert "../" not in res_data["fileName"]
    assert "evil_shell.docx" in res_data["fileName"]
