"""
Generates the 3 starter Word (.docx) templates along with schema.yaml, rules.yaml, and meta.yaml
matching the sample PDF files exactly:
  1. Room Permission Letter (EXEC MEET_public.pdf) -> templates/room_permission/
  2. Bill Certificate Form BC-R (Bill Certificate_public.pdf) -> templates/bill_certificate/
  3. Bill Summary (BIll_summary_public.pdf) -> templates/bill_summary/
"""

import os
from pathlib import Path
import yaml
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES_DIR = ROOT / "templates"
ASSETS_DIR = ROOT / "backend" / "assets"
SVNIT_LOGO = ASSETS_DIR / "svnit_logo.png"
ACM_LOGO = ASSETS_DIR / "acm_logo.png"


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_shading(cell, color_hex="F2F2F2"):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._element.get_or_add_tcPr().append(shd)


def set_cell_border(cell, **kwargs):
    """
    kwargs can be top, bottom, left, right.
    val: 'single', 'none', etc.
    sz: size in 1/8 pt (e.g. '4', '8')
    color: hex color string (e.g. '000000', 'CCCCCC')
    """
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            val = edge_data.get('val', 'single')
            color = edge_data.get('color', '000000')
            sz = str(edge_data.get('sz', '4'))
            b_el = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
            tcBorders.append(b_el)
        else:
            b_el = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_el)
    tcPr.append(tcBorders)


def add_word_header_and_footer(doc: Document, subtitle_text: str = "SARDAR VALLABHBHAI NATIONAL INSTITUTE OF TECHNOLOGY, SURAT"):
    """
    Adds a REAL Word header with left SVNIT logo, centered ACM/Institute title & subtitle,
    right ACM logo, and a footer with dynamic page numbering (PAGE of NUMPAGES) on every page.
    """
    for section in doc.sections:
        # Header setup
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.text = ""  # clear default paragraph

        tbl_hdr = header.add_table(rows=1, cols=3, width=Inches(6.8))
        tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_hdr.columns[0].width = Inches(1.1)
        tbl_hdr.columns[1].width = Inches(4.6)
        tbl_hdr.columns[2].width = Inches(1.1)

        # Left cell: SVNIT Logo
        cell_l = tbl_hdr.cell(0, 0)
        cell_l.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p_l = cell_l.paragraphs[0]
        p_l.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if SVNIT_LOGO.exists():
            p_l.add_run().add_picture(str(SVNIT_LOGO), width=Inches(0.85))

        # Center cell: ACM Blue Title & Subtitle
        cell_c = tbl_hdr.cell(0, 1)
        cell_c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p_c = cell_c.paragraphs[0]
        p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_title = p_c.add_run("ASSOCIATION for COMPUTING MACHINERY\n")
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(12)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(0x4F, 0x81, 0xBD)

        r_sub = p_c.add_run(subtitle_text)
        r_sub.font.name = "Times New Roman"
        r_sub.font.size = Pt(8.5)
        r_sub.font.bold = True
        r_sub.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

        # Right cell: ACM Logo
        cell_r = tbl_hdr.cell(0, 2)
        cell_r.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p_r = cell_r.paragraphs[0]
        p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if ACM_LOGO.exists():
            p_r.add_run().add_picture(str(ACM_LOGO), width=Inches(0.85))

        # Remove borders on header table
        for row in tbl_hdr.rows:
            for cell in row.cells:
                set_cell_border(cell)

        # Divider line paragraph under header
        p_line = header.add_paragraph()
        p_line.paragraph_format.space_before = Pt(4)
        p_line.paragraph_format.space_after = Pt(8)
        p_line_bdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="1" w:color="4F81BD"/></w:pBdr>')
        p_line._p.get_or_add_pPr().append(p_line_bdr)

        # Footer setup with page numbers: Page X of Y
        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ftr = p_ftr.add_run("Page ")
        r_ftr.font.name = "Times New Roman"
        r_ftr.font.size = Pt(9)
        r_ftr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        # Field: PAGE
        fld_page = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        p_ftr._p.append(fld_page)

        r_of = p_ftr.add_run(" of ")
        r_of.font.name = "Times New Roman"
        r_of.font.size = Pt(9)
        r_of.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        # Field: NUMPAGES
        fld_numpages = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
        p_ftr._p.append(fld_numpages)


# ==============================================================================
# 1. ROOM PERMISSION TEMPLATE
# ==============================================================================
def create_room_permission_template(out_dir: Path):
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    add_word_header_and_footer(doc, "SARDAR VALLABHBHAI NATIONAL INSTITUTE OF TECHNOLOGY, SURAT")

    # Reference No & Date
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(6)
    p_meta.paragraph_format.space_after = Pt(8)
    r_ref = p_meta.add_run("REF NO.: {{ ref_no }}")
    r_ref.font.name = "Times New Roman"
    r_ref.font.size = Pt(10)
    r_ref.font.bold = True

    p_meta.paragraph_format.tab_stops.add_tab_stop(Inches(6.8))
    r_tab = p_meta.add_run("\tDATE: {{ letter_date | date_dmy }}")
    r_tab.font.name = "Times New Roman"
    r_tab.font.size = Pt(10)
    r_tab.font.bold = True

    # Submitted to
    p_to = doc.add_paragraph()
    p_to.paragraph_format.space_after = Pt(6)
    r_to = p_to.add_run("SUBMITTED TO: {{ submitted_to }}")
    r_to.font.name = "Times New Roman"
    r_to.font.size = Pt(10)
    r_to.font.bold = True

    # Subject
    p_subj = doc.add_paragraph()
    p_subj.paragraph_format.space_after = Pt(6)
    r_subj = p_subj.add_run("SUBJECT: {{ subject }}")
    r_subj.font.name = "Times New Roman"
    r_subj.font.size = Pt(10)
    r_subj.font.bold = True

    # Schedule header
    p_sch = doc.add_paragraph()
    p_sch.paragraph_format.space_after = Pt(4)
    r_sch = p_sch.add_run("SCHEDULE OF EVENT:")
    r_sch.font.name = "Times New Roman"
    r_sch.font.size = Pt(10)
    r_sch.font.bold = True

    # Bordered "ABOUT THE EVENT" 2-column Table
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(5.0)

    # Row 0: Header "ABOUT THE EVENT"
    hdr_cell = table.cell(0, 0)
    hdr_cell.merge(table.cell(0, 1))
    set_cell_shading(hdr_cell, "4F81BD")
    set_cell_margins(hdr_cell, top=80, bottom=80, left=120, right=120)
    p_tbl_hdr = hdr_cell.paragraphs[0]
    p_tbl_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tbl_hdr = p_tbl_hdr.add_run("ABOUT THE EVENT")
    r_tbl_hdr.font.name = "Times New Roman"
    r_tbl_hdr.font.size = Pt(10.5)
    r_tbl_hdr.font.bold = True
    r_tbl_hdr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Row data
    rows_data = [
        ("Brief Description", "{{ description }}\n\nWe request permission to use {{ venue }} on {{ start_date | date_range(end_date) }} from {{ time_from | time_12h }} to {{ time_to | time_12h }} for conducting the meet."),
        ("Objective", "{{ objective }}"),
        ("Number of Participants", "{{ participants }}"),
        ("Mode", "{{ mode }}"),
        ("Event Organizers", "{{ organizers }}"),
    ]

    for idx, (label, val_expr) in enumerate(rows_data, start=1):
        cell_lbl = table.cell(idx, 0)
        cell_val = table.cell(idx, 1)

        set_cell_margins(cell_lbl, top=70, bottom=70, left=100, right=100)
        set_cell_margins(cell_val, top=70, bottom=70, left=100, right=100)
        set_cell_shading(cell_lbl, "F5F5F5")

        p_lbl = cell_lbl.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.size = Pt(9.5)
        r_lbl.font.bold = True

        p_val = cell_val.paragraphs[0]
        r_val = p_val.add_run(val_expr)
        r_val.font.name = "Times New Roman"
        r_val.font.size = Pt(9.5)

    # Apply solid borders to table
    border_kwargs = {'val': 'single', 'sz': '4', 'color': '888888'}
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top=border_kwargs, bottom=border_kwargs, left=border_kwargs, right=border_kwargs)

    # Spacing
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(8)
    p_sp.paragraph_format.space_after = Pt(6)

    # 2x2 Contact Grid (matching sample PDF)
    # Contact 1: Chairperson (Student)
    # Contact 2: Tech Assistant (Seminar Hall)
    # Contact 3: Faculty in-charge (Seminar Hall)
    # Contact 4: Faculty Chairman (ACM)
    contact_table = doc.add_table(rows=2, cols=2)
    contact_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    contact_table.columns[0].width = Inches(3.4)
    contact_table.columns[1].width = Inches(3.4)

    c_cells = [
        (contact_table.cell(0, 0), "contact_1"),
        (contact_table.cell(0, 1), "contact_2"),
        (contact_table.cell(1, 0), "contact_3"),
        (contact_table.cell(1, 1), "contact_4"),
    ]

    for cell, prefix in c_cells:
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        set_cell_border(cell)  # Borderless in sample
        p = cell.paragraphs[0]
        text_block = (
            f"Name: {{{{ {prefix}_name }}}}\n"
            f"Designation: {{{{ {prefix}_designation }}}}\n"
            f"{{% if {prefix}_admn %}}Admission No.: {{{{ {prefix}_admn }}}}\n{{% endif %}}"
            f"{{% if {prefix}_branch %}}Branch: {{{{ {prefix}_branch }}}}\n{{% endif %}}"
            f"{{% if {prefix}_mobile %}}Mobile No.: {{{{ {prefix}_mobile }}}}\n{{% endif %}}"
        )
        r = p.add_run(text_block)
        r.font.name = "Times New Roman"
        r.font.size = Pt(9.5)

    doc.save(str(out_dir / "template.docx"))

    # Meta
    meta = {
        "id": "room-permission",
        "name": "Room Permission Letter",
        "short_name": "Room Permission",
        "description": "Official letter addressed to the HoD requesting hall/lab booking with event details, schedule, and contacts.",
        "authority": "Head of Department (CSE Dept.)",
        "ref_category": "ROOM",
        "file_name": "room_permission_letter.docx",
        "version": 2,
    }
    (out_dir / "meta.yaml").write_text(yaml.dump(meta, sort_keys=False), encoding="utf-8")

    # Schema
    schema = {
        "placeholders": [
            {"key": "ref_no", "label": "Reference Number", "type": "text", "required": True, "question": "What is the dispatch reference number?", "never_ai": True, "masked": False, "shared": False, "section": "Administrative"},
            {"key": "letter_date", "label": "Date of Letter", "type": "date", "required": True, "question": "What is the date of the letter?", "never_ai": False, "masked": False, "shared": False, "section": "Administrative"},
            {"key": "submitted_to", "label": "Submitted To", "type": "text", "required": True, "question": "To whom is this request submitted?", "never_ai": False, "masked": False, "shared": False, "section": "General"},
            {"key": "subject", "label": "Subject", "type": "text", "required": True, "question": "What is the subject line?", "never_ai": False, "masked": False, "ai_draftable": True, "shared": False, "section": "General"},
            {"key": "event_title", "label": "Event Title", "type": "text", "required": True, "question": "What is the event title?", "never_ai": False, "masked": False, "shared": True, "section": "Event Details"},
            {"key": "venue", "label": "Venue / Room", "type": "text", "required": True, "question": "Which room or seminar hall is requested?", "never_ai": False, "masked": False, "shared": True, "section": "Event Details"},
            {"key": "start_date", "label": "Start Date", "type": "date", "required": True, "question": "What is the event start date?", "never_ai": False, "masked": False, "shared": True, "section": "Schedule"},
            {"key": "end_date", "label": "End Date", "type": "date", "required": True, "question": "What is the event end date?", "never_ai": False, "masked": False, "shared": True, "section": "Schedule"},
            {"key": "time_from", "label": "Start Time", "type": "text", "required": True, "question": "What time does the event begin?", "never_ai": False, "masked": False, "shared": True, "section": "Schedule"},
            {"key": "time_to", "label": "End Time", "type": "text", "required": True, "question": "What time does the event end?", "never_ai": False, "masked": False, "shared": True, "section": "Schedule"},
            {"key": "description", "label": "Brief Description", "type": "longtext", "required": True, "question": "Provide a brief description of the event.", "never_ai": False, "masked": False, "ai_draftable": True, "shared": False, "section": "Event Details"},
            {"key": "objective", "label": "Objective", "type": "longtext", "required": True, "question": "What is the primary objective of this event?", "never_ai": False, "masked": False, "ai_draftable": True, "shared": False, "section": "Event Details"},
            {"key": "participants", "label": "Expected Participants", "type": "text", "required": True, "question": "How many participants are expected?", "never_ai": False, "masked": False, "shared": True, "section": "Event Details"},
            {"key": "mode", "label": "Mode of Event", "type": "select", "required": True, "question": "Is the event offline, online, or hybrid?", "options": ["Offline", "Online", "Hybrid"], "never_ai": False, "masked": False, "shared": True, "section": "Event Details"},
            {"key": "organizers", "label": "Event Organizers", "type": "text", "required": True, "question": "Who is organizing this event?", "never_ai": False, "masked": False, "shared": True, "section": "Event Details"},
            {"key": "faculty_coordinator", "label": "Faculty Coordinator", "type": "text", "required": True, "question": "Who is the faculty coordinator?", "never_ai": False, "masked": False, "shared": True, "section": "Signatories"},
            {"key": "contact_1_name", "label": "Student Head Name", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Contacts"},
            {"key": "contact_1_designation", "label": "Student Head Designation", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Contacts"},
            {"key": "contact_1_admn", "label": "Student Head Admission No.", "type": "text", "required": False, "never_ai": False, "masked": False, "section": "Contacts"},
            {"key": "contact_1_branch", "label": "Student Head Branch", "type": "text", "required": False, "never_ai": False, "masked": False, "section": "Contacts"},
            {"key": "contact_1_mobile", "label": "Student Head Mobile No.", "type": "text", "required": False, "never_ai": True, "masked": True, "section": "Contacts"},
            {"key": "contact_2_name", "label": "Tech Assistant Name", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Contacts"},
            {"key": "contact_2_designation", "label": "Tech Assistant Designation", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Contacts"},
            {"key": "contact_2_branch", "label": "Tech Assistant Branch", "type": "text", "required": False, "never_ai": False, "masked": False, "section": "Contacts"},
            {"key": "contact_2_mobile", "label": "Tech Assistant Mobile No.", "type": "text", "required": False, "never_ai": True, "masked": True, "section": "Contacts"},
            {"key": "contact_3_name", "label": "Faculty In-charge Name", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Contacts"},
            {"key": "contact_3_designation", "label": "Faculty In-charge Designation", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Contacts"},
            {"key": "contact_3_branch", "label": "Faculty In-charge Branch", "type": "text", "required": False, "never_ai": False, "masked": False, "section": "Contacts"},
            {"key": "contact_4_name", "label": "Faculty Chairman Name", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Contacts"},
            {"key": "contact_4_designation", "label": "Faculty Chairman Designation", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Contacts"},
            {"key": "contact_4_branch", "label": "Faculty Chairman Branch", "type": "text", "required": False, "never_ai": False, "masked": False, "section": "Contacts"},
        ]
    }
    (out_dir / "schema.yaml").write_text(yaml.dump(schema, sort_keys=False), encoding="utf-8")

    # Rules
    rules = {
        "rules": [
            {
                "id": "room-offline",
                "description": "Recommended when the event requires an on-campus hall or seminar room.",
                "condition": "mode != 'Online' and venue != None",
            }
        ]
    }
    (out_dir / "rules.yaml").write_text(yaml.dump(rules, sort_keys=False), encoding="utf-8")


# ==============================================================================
# 2. BILL CERTIFICATE (FORM BC-R) TEMPLATE
# ==============================================================================
def create_bill_certificate_template(out_dir: Path):
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    add_word_header_and_footer(doc, "SARDAR VALLABHBHAI NATIONAL INSTITUTE OF TECHNOLOGY, SURAT")

    # Form BC-R Title
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_top.paragraph_format.space_before = Pt(4)
    p_top.paragraph_format.space_after = Pt(2)
    r_dsw = p_top.add_run("DEAN STUDENT WELFARE\nFORM BC-R\nBILL CERTIFICATE (RECURRING EXPENDITURE)")
    r_dsw.font.name = "Times New Roman"
    r_dsw.font.size = Pt(10.5)
    r_dsw.font.bold = True

    # Details Block
    p_dtl = doc.add_paragraph()
    p_dtl.paragraph_format.space_before = Pt(4)
    p_dtl.paragraph_format.space_after = Pt(4)
    r_dtl = p_dtl.add_run(
        "Chapter Name: {{ chapter_name }}\t\tCode No.: {{ chapter_code }}\n"
        "Approval Note No.: {{ approval_note_no }}\t\tDate: {{ approval_date | date_dmy }}\n"
        "Original Approval Note/Comparative Statement attached herewith/sent with bill No. {{ bill_no }} dtd.: {{ bill_date | date_dmy }} for Rs. {{ bill_amount }} of M/s. {{ supplier_name }}\n"
        "Purchase Order No. and date : {{ purchase_order }}\n"
        "Head of Account (Code No.) : {{ head_of_account }}\n"
        "Advance drawn in Rs. (if any) : {{ advance_drawn }}"
    )
    r_dtl.font.name = "Times New Roman"
    r_dtl.font.size = Pt(9.5)

    # Certification statement
    p_cert = doc.add_paragraph()
    p_cert.paragraph_format.space_before = Pt(6)
    p_cert.paragraph_format.space_after = Pt(6)
    r_cert = p_cert.add_run(
        '"I, {{ certifying_person }}, personally satisfied that the goods (described below) purchased are of the requisite quality and specification and have been purchased from a reliable supplier / contractor at a reasonable price."'
    )
    r_cert.font.name = "Times New Roman"
    r_cert.font.size = Pt(9.5)
    r_cert.font.italic = True
    r_cert.font.bold = True

    # Items Table using {%tr for item in items %}
    tbl_items = doc.add_table(rows=6, cols=5)
    tbl_items.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_items.columns[0].width = Inches(0.6)
    tbl_items.columns[1].width = Inches(3.2)
    tbl_items.columns[2].width = Inches(0.8)
    tbl_items.columns[3].width = Inches(1.1)
    tbl_items.columns[4].width = Inches(1.2)

    # Header Row
    headers = ["Sr. No.", "Name of the Equipment/ item(s)", "Quantity", "Unit Cost (Rs.)", "Total Cost of Item(s) (Rs.)"]
    for i, h in enumerate(headers):
        cell = tbl_items.cell(0, i)
        set_cell_shading(cell, "EAEAEA")
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.size = Pt(9)
        r.font.bold = True

    # Row 1: Loop start
    tbl_items.cell(1, 0).paragraphs[0].add_run("{%tr for item in items %}")

    # Row 2: Data row
    row_data = tbl_items.rows[2]
    row_data.cells[0].paragraphs[0].add_run("{{ loop.index }}")
    row_data.cells[1].paragraphs[0].add_run("{{ item.item_name }}")
    row_data.cells[2].paragraphs[0].add_run("{{ item.qty }}")
    row_data.cells[3].paragraphs[0].add_run("{{ item.unit_cost }}")
    row_data.cells[4].paragraphs[0].add_run("{{ item.total }}")

    # Row 3: Loop end
    tbl_items.cell(3, 0).paragraphs[0].add_run("{%tr endfor %}")

    # Total Row
    row_tot = tbl_items.rows[4]
    cell_tot_lbl = row_tot.cells[0]
    cell_tot_lbl.merge(row_tot.cells[3])
    cell_tot_lbl.paragraphs[0].add_run("Total Amount (Rs.)").font.bold = True
    row_tot.cells[4].paragraphs[0].add_run("{{ total_amount }}").font.bold = True

    # Grand Total Row
    row_gt = tbl_items.rows[5]
    cell_gt_lbl = row_gt.cells[0]
    cell_gt_lbl.merge(row_gt.cells[3])
    cell_gt_lbl.paragraphs[0].add_run("Total Estimated Amount Including Taxes (Rs.)").font.bold = True
    row_gt.cells[4].paragraphs[0].add_run("{{ grand_total }}").font.bold = True

    bdr = {'val': 'single', 'sz': '4', 'color': '888888'}
    for row in tbl_items.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            set_cell_border(cell, top=bdr, bottom=bdr, left=bdr, right=bdr)

    # SUPPLIER DETAILS
    p_sup_hdr = doc.add_paragraph()
    p_sup_hdr.paragraph_format.space_before = Pt(8)
    p_sup_hdr.paragraph_format.space_after = Pt(2)
    p_sup_hdr.add_run("SUPPLIER DETAILS").font.bold = True

    tbl_sup = doc.add_table(rows=5, cols=2)
    tbl_sup.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_sup.columns[0].width = Inches(2.2)
    tbl_sup.columns[1].width = Inches(4.7)

    sup_fields = [
        ("Name of the Supplier", "{{ supplier_name }}"),
        ("GST No.", "{{ supplier_gst }}"),
        ("Any other Taxes", "{{ other_taxes }}"),
        ("Invoice No. and date", "{{ invoice_no }} and {{ invoice_date | date_dmy }}"),
        ("Amount (Rs.)", "{{ grand_total }}"),
    ]
    for idx, (lbl, val) in enumerate(sup_fields):
        c0 = tbl_sup.cell(idx, 0)
        c1 = tbl_sup.cell(idx, 1)
        set_cell_margins(c0, top=40, bottom=40, left=80, right=80)
        set_cell_margins(c1, top=40, bottom=40, left=80, right=80)
        set_cell_border(c0, top=bdr, bottom=bdr, left=bdr, right=bdr)
        set_cell_border(c1, top=bdr, bottom=bdr, left=bdr, right=bdr)
        c0.paragraphs[0].add_run(lbl).font.size = Pt(9)
        c1.paragraphs[0].add_run(val).font.size = Pt(9)

    # PAYMENT DETAILS
    p_pay_hdr = doc.add_paragraph()
    p_pay_hdr.paragraph_format.space_before = Pt(8)
    p_pay_hdr.paragraph_format.space_after = Pt(2)
    p_pay_hdr.add_run("PAYMENT DETAILS").font.bold = True

    tbl_pay = doc.add_table(rows=6, cols=2)
    tbl_pay.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pay.columns[0].width = Inches(2.2)
    tbl_pay.columns[1].width = Inches(4.7)

    pay_fields = [
        ("Name of the Party", "{{ payee_name }}"),
        ("Account Number", "{{ account_number }}"),
        ("Account Holder Name", "{{ account_holder }}"),
        ("Name of the Bank", "{{ bank_name }}"),
        ("Name of the Branch", "{{ bank_branch }}"),
        ("IFSC Code", "{{ ifsc }}"),
    ]
    for idx, (lbl, val) in enumerate(pay_fields):
        c0 = tbl_pay.cell(idx, 0)
        c1 = tbl_pay.cell(idx, 1)
        set_cell_margins(c0, top=40, bottom=40, left=80, right=80)
        set_cell_margins(c1, top=40, bottom=40, left=80, right=80)
        set_cell_border(c0, top=bdr, bottom=bdr, left=bdr, right=bdr)
        set_cell_border(c1, top=bdr, bottom=bdr, left=bdr, right=bdr)
        c0.paragraphs[0].add_run(lbl).font.size = Pt(9)
        c1.paragraphs[0].add_run(val).font.size = Pt(9)

    # PAGE 2 BREAK
    doc.add_page_break()

    # STOCK ENTRY DETAILS
    p_stk = doc.add_paragraph()
    p_stk.paragraph_format.space_before = Pt(4)
    p_stk.paragraph_format.space_after = Pt(2)
    p_stk.add_run("STOCK ENTRY DETAILS").font.bold = True
    p_stk_sub = doc.add_paragraph()
    p_stk_sub.paragraph_format.space_after = Pt(4)
    p_stk_sub.add_run("Name of the Laboratory / Office: {{ stock_office }}").font.size = Pt(9.5)

    tbl_stk = doc.add_table(rows=5, cols=4)
    tbl_stk.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_stk.columns[0].width = Inches(1.8)
    tbl_stk.columns[1].width = Inches(1.8)
    tbl_stk.columns[2].width = Inches(1.8)
    tbl_stk.columns[3].width = Inches(1.5)

    stk_headers = ["Recurring Register", "Consumable Register", "Miscellaneous Register", "Clerk / Staff Signature"]
    for i, h in enumerate(stk_headers):
        c = tbl_stk.cell(0, i)
        set_cell_shading(c, "EAEAEA")
        set_cell_margins(c, top=40, bottom=40, left=60, right=60)
        c.paragraphs[0].add_run(h).font.bold = True

    stk_rows = [
        ("Name/Type: {{ stock_register_type }}", "-", "-", ""),
        ("Serial Number: {{ stock_serial_no }}", "-", "-", ""),
        ("Page Number: {{ stock_page_no }}", "-", "-", ""),
        ("Date: {{ stock_date | date_dmy }}", "-", "-", ""),
    ]
    for r_idx, (c0, c1, c2, c3) in enumerate(stk_rows, start=1):
        for col_idx, text in enumerate([c0, c1, c2, c3]):
            cell = tbl_stk.cell(r_idx, col_idx)
            set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
            cell.paragraphs[0].add_run(text).font.size = Pt(8.5)

    for row in tbl_stk.rows:
        for cell in row.cells:
            set_cell_border(cell, top=bdr, bottom=bdr, left=bdr, right=bdr)

    # INDENTER DETAILS
    p_ind = doc.add_paragraph()
    p_ind.paragraph_format.space_before = Pt(8)
    p_ind.paragraph_format.space_after = Pt(2)
    p_ind.add_run("INDENTER DETAILS").font.bold = True

    tbl_ind = doc.add_table(rows=2, cols=2)
    tbl_ind.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ind.columns[0].width = Inches(3.5)
    tbl_ind.columns[1].width = Inches(3.4)

    c00 = tbl_ind.cell(0, 0)
    c01 = tbl_ind.cell(0, 1)
    c10 = tbl_ind.cell(1, 0)
    c11 = tbl_ind.cell(1, 1)

    c00.paragraphs[0].add_run("Name of Faculty Chairman:\n{{ faculty_chairman_name }}").font.size = Pt(9.5)
    c01.paragraphs[0].add_run("Designation:\n{{ faculty_chairman_designation }}").font.size = Pt(9.5)
    c10.paragraphs[0].add_run("Department: {{ department }}").font.size = Pt(9.5)
    c11.paragraphs[0].add_run("Signature with Date:\n\n_______________________").font.size = Pt(9.5)

    for row in tbl_ind.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            set_cell_border(cell, top=bdr, bottom=bdr, left=bdr, right=bdr)

    # Recommendation
    p_rec = doc.add_paragraph()
    p_rec.paragraph_format.space_before = Pt(8)
    p_rec.paragraph_format.space_after = Pt(4)
    p_rec.add_run("Recommendation (Head of the Department)\nRecommended Amount: Rs. {{ grand_total }}\nApproved and Recommended for Bill Payment\n\nDate: {{ approval_date | date_dmy }}\t\t\t\tDean Student Welfare").font.size = Pt(9.5)

    # (FOR USE IN ACCOUNT SECTION) - Left strictly blank as per instructions
    p_acc = doc.add_paragraph()
    p_acc.paragraph_format.space_before = Pt(8)
    p_acc.paragraph_format.space_after = Pt(4)
    r_acc_hdr = p_acc.add_run("(FOR USE IN ACCOUNT SECTION)")
    r_acc_hdr.font.bold = True
    r_acc_hdr.font.size = Pt(10)

    p_acc_body = doc.add_paragraph()
    p_acc_body.add_run(
        "Date:\n"
        "Head of Account :\n"
        "Budget Code No. :\n"
        "Mode of Payment  : Cash / Cheque / Online / Adjusted Against Advance\n"
        "Details of Mode of Payment :\n"
        "Voucher No. and Date :\n"
        "Passed for payment/adjustment of Amount in (Rs.) _________________ (in figs.) ________________\n"
        "(in words) _________________________________________________________________\n\n\n"
        "Office Supt.\t\t\tDy. Registrar (A/cs.)\t\t\tRegistrar / Director"
    ).font.size = Pt(9)

    doc.save(str(out_dir / "template.docx"))

    # Meta
    meta = {
        "id": "bill-certificate",
        "name": "Bill Certificate (Form BC-R)",
        "short_name": "Bill Certificate",
        "description": "Institute Form BC-R for recurring club expenses with itemized table, GST, payee details, and stock entry.",
        "authority": "Dean (Student Welfare)",
        "ref_category": "BILL",
        "file_name": "bill_certificate_form_bcr.docx",
        "version": 2,
    }
    (out_dir / "meta.yaml").write_text(yaml.dump(meta, sort_keys=False), encoding="utf-8")

    # Schema
    schema = {
        "placeholders": [
            {"key": "chapter_name", "label": "Chapter Name", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "General"},
            {"key": "chapter_code", "label": "Chapter Code No.", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "General"},
            {"key": "approval_note_no", "label": "Approval Note No.", "type": "text", "required": True, "never_ai": True, "masked": False, "section": "Approval Details"},
            {"key": "approval_date", "label": "Approval Date", "type": "date", "required": True, "never_ai": True, "masked": False, "section": "Approval Details"},
            {"key": "bill_no", "label": "Bill No.", "type": "text", "required": True, "never_ai": True, "masked": False, "section": "Invoice"},
            {"key": "bill_date", "label": "Bill Date", "type": "date", "required": True, "never_ai": True, "masked": False, "section": "Invoice"},
            {"key": "bill_amount", "label": "Bill Amount (Rs.)", "type": "number", "required": True, "never_ai": True, "masked": False, "section": "Invoice"},
            {"key": "supplier_name", "label": "Supplier Name", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Supplier Details"},
            {"key": "purchase_order", "label": "Purchase Order No. & Date", "type": "text", "required": True, "never_ai": True, "masked": False, "section": "Approval Details"},
            {"key": "head_of_account", "label": "Head of Account (Code No.)", "type": "text", "required": True, "never_ai": True, "masked": False, "section": "Approval Details"},
            {"key": "advance_drawn", "label": "Advance Drawn in Rs.", "type": "text", "required": True, "never_ai": True, "masked": False, "section": "Approval Details"},
            {"key": "certifying_person", "label": "Certifying Person Name", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Certification"},
            {"key": "items", "label": "Purchased Items", "type": "table", "required": True, "never_ai": True, "masked": False, "section": "Itemized Expenses", "columns": [
                {"key": "sr_no", "label": "Sr No.", "type": "number"},
                {"key": "item_name", "label": "Item Name", "type": "text"},
                {"key": "qty", "label": "Qty", "type": "number"},
                {"key": "unit_cost", "label": "Unit Cost (Rs.)", "type": "number"},
                {"key": "total", "label": "Total Cost (Rs.)", "type": "number"},
            ]},
            {"key": "total_amount", "label": "Total Amount (Rs.)", "type": "number", "required": False, "never_ai": True, "masked": False, "section": "Itemized Expenses"},
            {"key": "taxes", "label": "Taxes / Other Charges (Rs.)", "type": "number", "required": False, "never_ai": True, "masked": False, "section": "Itemized Expenses"},
            {"key": "grand_total", "label": "Grand Total Amount (Rs.)", "type": "number", "required": False, "never_ai": True, "masked": False, "section": "Itemized Expenses"},
            {"key": "supplier_gst", "label": "Supplier GST No.", "type": "text", "required": True, "never_ai": True, "masked": True, "section": "Supplier Details"},
            {"key": "other_taxes", "label": "Any Other Taxes", "type": "text", "required": False, "never_ai": False, "masked": False, "section": "Supplier Details"},
            {"key": "invoice_no", "label": "Invoice Number", "type": "text", "required": True, "never_ai": True, "masked": False, "section": "Supplier Details"},
            {"key": "invoice_date", "label": "Invoice Date", "type": "date", "required": True, "never_ai": True, "masked": False, "section": "Supplier Details"},
            {"key": "payee_name", "label": "Party / Payee Name", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Payment Details"},
            {"key": "account_number", "label": "Bank Account Number", "type": "text", "required": True, "never_ai": True, "masked": True, "section": "Payment Details"},
            {"key": "account_holder", "label": "Account Holder Name", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Payment Details"},
            {"key": "bank_name", "label": "Bank Name", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Payment Details"},
            {"key": "bank_branch", "label": "Bank Branch", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Payment Details"},
            {"key": "ifsc", "label": "IFSC Code", "type": "text", "required": True, "never_ai": True, "masked": True, "section": "Payment Details"},
            {"key": "stock_office", "label": "Laboratory / Office Name", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Stock Entry"},
            {"key": "stock_register_type", "label": "Register Name/Type", "type": "text", "required": False, "never_ai": False, "masked": False, "section": "Stock Entry"},
            {"key": "stock_serial_no", "label": "Stock Serial Number", "type": "text", "required": False, "never_ai": False, "masked": False, "section": "Stock Entry"},
            {"key": "stock_page_no", "label": "Stock Page Number", "type": "text", "required": False, "never_ai": False, "masked": False, "section": "Stock Entry"},
            {"key": "stock_date", "label": "Stock Entry Date", "type": "date", "required": False, "never_ai": False, "masked": False, "section": "Stock Entry"},
            {"key": "faculty_chairman_name", "label": "Faculty Chairman Name", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Indenter Details"},
            {"key": "faculty_chairman_designation", "label": "Faculty Chairman Designation", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Indenter Details"},
            {"key": "department", "label": "Department", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Indenter Details"},
        ]
    }
    (out_dir / "schema.yaml").write_text(yaml.dump(schema, sort_keys=False), encoding="utf-8")

    # Rules
    rules = {
        "rules": [
            {
                "id": "bill-expenses",
                "description": "Recommended whenever equipment, refreshments, printing, or club items are purchased.",
                "condition": "hasExpenses == True",
            }
        ]
    }
    (out_dir / "rules.yaml").write_text(yaml.dump(rules, sort_keys=False), encoding="utf-8")


# ==============================================================================
# 3. BILL SUMMARY TEMPLATE
# ==============================================================================
def create_bill_summary_template(out_dir: Path):
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    add_word_header_and_footer(doc, "SARDAR VALLABHBHAI NATIONAL INSTITUTE OF TECHNOLOGY, SURAT")

    # Date
    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_date.paragraph_format.space_before = Pt(6)
    p_date.paragraph_format.space_after = Pt(6)
    p_date.add_run("Date: {{ letter_date | date_dmy }}").font.bold = True

    # Title & Subtitle
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    r_tit = p_title.add_run("BILL SUMMARY")
    r_tit.font.name = "Times New Roman"
    r_tit.font.size = Pt(13)
    r_tit.font.bold = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(10)
    r_sub = p_sub.add_run("Reimbursement for Cost incurred in {{ event_title }}")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(10.5)
    r_sub.font.bold = True

    # Table 1: Reimbursements using {%tr for row in reimbursements %}
    tbl_reimb = doc.add_table(rows=5, cols=7)
    tbl_reimb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_reimb.columns[0].width = Inches(0.5)
    tbl_reimb.columns[1].width = Inches(1.5)
    tbl_reimb.columns[2].width = Inches(1.5)
    tbl_reimb.columns[3].width = Inches(1.2)
    tbl_reimb.columns[4].width = Inches(0.8)
    tbl_reimb.columns[5].width = Inches(0.7)
    tbl_reimb.columns[6].width = Inches(0.8)

    headers = ["Sr. No.", "Name", "Account Details", "To Be Paid to", "Prize Money (INR)", "Other", "Total (INR)"]
    for i, h in enumerate(headers):
        cell = tbl_reimb.cell(0, i)
        set_cell_shading(cell, "EAEAEA")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.size = Pt(8.5)
        r.font.bold = True

    # Row 1: Loop start
    tbl_reimb.cell(1, 0).paragraphs[0].add_run("{%tr for row in reimbursements %}")

    # Row 2: Data row
    r2 = tbl_reimb.rows[2]
    r2.cells[0].paragraphs[0].add_run("{{ loop.index }}")
    r2.cells[1].paragraphs[0].add_run("{{ row.name }}")
    r2.cells[2].paragraphs[0].add_run("{{ row.account_details }}")
    r2.cells[3].paragraphs[0].add_run("{{ row.paid_to }}")
    r2.cells[4].paragraphs[0].add_run("{{ row.prize_money }}")
    r2.cells[5].paragraphs[0].add_run("{{ row.other }}")
    r2.cells[6].paragraphs[0].add_run("{{ row.total }}")

    # Row 3: Loop end
    tbl_reimb.cell(3, 0).paragraphs[0].add_run("{%tr endfor %}")

    # Row 4: Grand Total
    r4 = tbl_reimb.rows[4]
    tot_lbl = r4.cells[0]
    tot_lbl.merge(r4.cells[3])
    tot_lbl.paragraphs[0].add_run("Grand Total").font.bold = True
    r4.cells[4].paragraphs[0].add_run("{{ total_prizes }}").font.bold = True
    r4.cells[5].paragraphs[0].add_run("{{ total_other }}").font.bold = True
    r4.cells[6].paragraphs[0].add_run("{{ grand_total }}").font.bold = True

    bdr = {'val': 'single', 'sz': '4', 'color': '888888'}
    for row in tbl_reimb.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
            set_cell_border(cell, top=bdr, bottom=bdr, left=bdr, right=bdr)

    # Spacing
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(20)

    # Signatures Block (3 columns from club profile)
    tbl_sig = doc.add_table(rows=1, cols=3)
    tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_sig.columns[0].width = Inches(2.3)
    tbl_sig.columns[1].width = Inches(2.3)
    tbl_sig.columns[2].width = Inches(2.3)

    sig_defs = [
        ("Student Head", "{{ student_chairperson_name }}", "ACM NIT-Surat"),
        ("Student Secretary", "{{ student_secretary_name }}", "ACM NIT-Surat"),
        ("Faculty Chairman", "{{ faculty_chairman_name }}", "ACM NIT-Surat"),
    ]
    for idx, (role, name, org) in enumerate(sig_defs):
        c = tbl_sig.cell(0, idx)
        set_cell_margins(c, top=60, bottom=60, left=60, right=60)
        set_cell_border(c)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"{role}\n{name}\n{org}").font.size = Pt(9.5)

    # PAGE 2 BREAK
    doc.add_page_break()

    # Table 2: Other Bank Account Details
    p_oth_hdr = doc.add_paragraph()
    p_oth_hdr.paragraph_format.space_before = Pt(6)
    p_oth_hdr.paragraph_format.space_after = Pt(8)
    p_oth_hdr.add_run("Other Bank Account Details").font.bold = True

    tbl_oth = doc.add_table(rows=5, cols=5)
    tbl_oth.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_oth.columns[0].width = Inches(0.6)
    tbl_oth.columns[1].width = Inches(2.0)
    tbl_oth.columns[2].width = Inches(2.0)
    tbl_oth.columns[3].width = Inches(1.2)
    tbl_oth.columns[4].width = Inches(1.2)

    oth_headers = ["Sr. No.", "Name", "Account No.", "IFSC Code", "Amount (Rs.)"]
    for i, h in enumerate(oth_headers):
        cell = tbl_oth.cell(0, i)
        set_cell_shading(cell, "EAEAEA")
        cell.paragraphs[0].add_run(h).font.bold = True

    # Row 1: Loop start
    tbl_oth.cell(1, 0).paragraphs[0].add_run("{%tr for row in other_accounts %}")

    # Row 2: Data row
    r2_oth = tbl_oth.rows[2]
    r2_oth.cells[0].paragraphs[0].add_run("{{ loop.index }}")
    r2_oth.cells[1].paragraphs[0].add_run("{{ row.name }}")
    r2_oth.cells[2].paragraphs[0].add_run("{{ row.account_no }}")
    r2_oth.cells[3].paragraphs[0].add_run("{{ row.ifsc }}")
    r2_oth.cells[4].paragraphs[0].add_run("{{ row.amount }}")

    # Row 3: Loop end
    tbl_oth.cell(3, 0).paragraphs[0].add_run("{%tr endfor %}")

    # Row 4: Total
    r4_oth = tbl_oth.rows[4]
    cell_tot = r4_oth.cells[0]
    cell_tot.merge(r4_oth.cells[3])
    cell_tot.paragraphs[0].add_run("Total:").font.bold = True
    r4_oth.cells[4].paragraphs[0].add_run("{{ other_accounts_total }}").font.bold = True

    for row in tbl_oth.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
            set_cell_border(cell, top=bdr, bottom=bdr, left=bdr, right=bdr)

    # Single signature block for page 2
    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(25)

    p_sig2 = doc.add_paragraph()
    p_sig2.paragraph_format.space_before = Pt(10)
    p_sig2.add_run("Faculty Chairman\n{{ faculty_chairman_name }}\nACM NIT-Surat").font.size = Pt(9.5)

    doc.save(str(out_dir / "template.docx"))

    # Meta
    meta = {
        "id": "bill-summary",
        "name": "Bill Summary & Reimbursement Sheet",
        "short_name": "Bill Summary",
        "description": "Consolidated reimbursement summary with payment breakdown, other accounts table, and club signatories.",
        "authority": "Dean (Student Welfare)",
        "ref_category": "BILL",
        "file_name": "bill_summary.docx",
        "version": 2,
    }
    (out_dir / "meta.yaml").write_text(yaml.dump(meta, sort_keys=False), encoding="utf-8")

    # Schema
    schema = {
        "placeholders": [
            {"key": "letter_date", "label": "Date of Summary", "type": "date", "required": True, "never_ai": False, "masked": False, "section": "General"},
            {"key": "event_title", "label": "Event Title", "type": "text", "required": True, "never_ai": False, "masked": False, "shared": True, "section": "General"},
            {"key": "reimbursements", "label": "Reimbursement Items", "type": "table", "required": True, "never_ai": True, "masked": False, "section": "Reimbursements", "columns": [
                {"key": "sr_no", "label": "Sr No.", "type": "number"},
                {"key": "name", "label": "Payee Name", "type": "text"},
                {"key": "account_details", "label": "Account Details", "type": "text"},
                {"key": "paid_to", "label": "To Be Paid To", "type": "text"},
                {"key": "prize_money", "label": "Prize Money (INR)", "type": "number"},
                {"key": "other", "label": "Other Expenses (INR)", "type": "number"},
                {"key": "total", "label": "Total Amount (INR)", "type": "number"},
            ]},
            {"key": "total_prizes", "label": "Total Prize Money (INR)", "type": "number", "required": False, "never_ai": True, "masked": False, "section": "Reimbursements"},
            {"key": "total_other", "label": "Total Other Expenses (INR)", "type": "number", "required": False, "never_ai": True, "masked": False, "section": "Reimbursements"},
            {"key": "grand_total", "label": "Grand Total (INR)", "type": "number", "required": False, "never_ai": True, "masked": False, "section": "Reimbursements"},
            {"key": "other_accounts", "label": "Other Bank Accounts", "type": "table", "required": False, "never_ai": True, "masked": False, "section": "Other Accounts", "columns": [
                {"key": "sr_no", "label": "Sr No.", "type": "number"},
                {"key": "name", "label": "Account Name", "type": "text"},
                {"key": "account_no", "label": "Account Number", "type": "text"},
                {"key": "ifsc", "label": "IFSC Code", "type": "text"},
                {"key": "amount", "label": "Amount (INR)", "type": "number"},
            ]},
            {"key": "other_accounts_total", "label": "Other Accounts Total (INR)", "type": "number", "required": False, "never_ai": True, "masked": False, "section": "Other Accounts"},
            {"key": "student_chairperson_name", "label": "Student Head Name", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Signatories"},
            {"key": "student_secretary_name", "label": "Student Secretary Name", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Signatories"},
            {"key": "faculty_chairman_name", "label": "Faculty Chairman Name", "type": "text", "required": True, "never_ai": False, "masked": False, "section": "Signatories"},
        ]
    }
    (out_dir / "schema.yaml").write_text(yaml.dump(schema, sort_keys=False), encoding="utf-8")

    # Rules
    rules = {
        "rules": [
            {
                "id": "summary-expenses-or-prizes",
                "description": "Recommended whenever reimbursements, prizes, or multiple payments must be cleared.",
                "condition": "hasExpenses == True or hasPrizes == True",
            }
        ]
    }
    (out_dir / "rules.yaml").write_text(yaml.dump(rules, sort_keys=False), encoding="utf-8")


def main():
    print("Building templates in:", TEMPLATES_DIR)
    create_room_permission_template(TEMPLATES_DIR / "room_permission")
    create_bill_certificate_template(TEMPLATES_DIR / "bill_certificate")
    create_bill_summary_template(TEMPLATES_DIR / "bill_summary")
    print("Successfully built all 3 templates with Word headers, logos, and real row tables.")


if __name__ == "__main__":
    main()
